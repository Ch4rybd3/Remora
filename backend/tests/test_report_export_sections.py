"""
A report section reaches both exporters.

The two export paths disagreed. `_render_markdown` looped over the case's
`report_sections_data` and substituted every slug it found. `_render_docx`
scanned each paragraph against a *fixed* set of block tag names, so a section
slug matched nothing, fell through to the text pass - which had no value for it
either, since the context holds registered tags only - and was written into the
delivered Word document as a literal `{{slug}}`.

The DOCX branch meant to handle it was unreachable for exactly that reason: it
sat behind a `block_found` that could never be a section name.

These tests export the same template in both formats and assert they agree.
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient

SCOPE_TEXT = "Three file servers and one domain controller."

TEMPLATE_BODY = [
    "{{case.title}}",
    "{{scope}}",
    "{{containment}}",
    "{{ioc_table}}",
]


@pytest.fixture()
def case_with_sections(db_session) -> str:
    import uuid as _uuid

    from app.models.case import Case

    new_id = str(_uuid.uuid4())
    db_session.add(Case(
        id=new_id,
        title="Acme ransomware",
        report_sections_data=json.dumps({
            "scope":       SCOPE_TEXT,
            "containment": "",          # created but not written yet
        }),
    ))
    db_session.commit()
    return new_id


def _upload(client: TestClient, name: str, filename: str, body: bytes, mime: str) -> int:
    response = client.post(
        "/api/v1/report-doc-templates/upload",
        data={"name": name, "description": ""},
        files={"file": (filename, body, mime)},
    )
    assert response.status_code == 200, response.text
    return response.json()["id"]


@pytest.fixture()
def markdown_template(auth_client: TestClient) -> int:
    body = "\n\n".join(TEMPLATE_BODY).encode()
    return _upload(auth_client, "MD sections", "sections.md", body, "text/markdown")


@pytest.fixture()
def docx_template(auth_client: TestClient) -> int:
    from docx import Document

    doc = Document()
    for line in TEMPLATE_BODY:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)

    return _upload(
        auth_client, "DOCX sections", "sections.docx", buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def _generate(client: TestClient, template_id: int, case_id: str) -> bytes:
    response = client.post(
        f"/api/v1/report-doc-templates/{template_id}/generate/{case_id}")
    assert response.status_code == 200, response.text
    return response.content


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ─── Markdown, which already worked ───────────────────────────────────────────

def test_markdown_renders_a_written_section(auth_client, markdown_template, case_with_sections):
    text = _generate(auth_client, markdown_template, case_with_sections).decode()

    assert SCOPE_TEXT in text
    assert "{{scope}}" not in text


def test_markdown_marks_an_unwritten_section(auth_client, markdown_template, case_with_sections):
    text = _generate(auth_client, markdown_template, case_with_sections).decode()

    assert "Section 'containment' not written" in text
    assert "{{containment}}" not in text


# ─── DOCX, which did not ──────────────────────────────────────────────────────

def test_docx_renders_a_written_section(auth_client, docx_template, case_with_sections):
    """The regression. This used to emit the literal text `{{scope}}`."""
    text = _docx_text(_generate(auth_client, docx_template, case_with_sections))

    assert SCOPE_TEXT in text
    assert "{{scope}}" not in text


def test_docx_marks_an_unwritten_section(auth_client, docx_template, case_with_sections):
    text = _docx_text(_generate(auth_client, docx_template, case_with_sections))

    assert "Section 'containment' not written" in text
    assert "{{containment}}" not in text


# ─── Neither path leaves a tag behind ─────────────────────────────────────────

def test_no_tag_survives_either_export(
    auth_client, markdown_template, docx_template, case_with_sections,
):
    """
    The whole class of bug, asserted once: whatever the template asked for,
    nothing shaped like `{{...}}` may reach a client-facing document.
    """
    md   = _generate(auth_client, markdown_template, case_with_sections).decode()
    docx = _docx_text(_generate(auth_client, docx_template, case_with_sections))

    for rendered, fmt in ((md, "markdown"), (docx, "docx")):
        assert "{{" not in rendered, f"an unresolved tag reached the {fmt} export"


def test_registered_tags_still_resolve_in_both(
    auth_client, markdown_template, docx_template, case_with_sections,
):
    """Sections must not have displaced the built-ins on the way in."""
    md   = _generate(auth_client, markdown_template, case_with_sections).decode()
    docx = _docx_text(_generate(auth_client, docx_template, case_with_sections))

    for rendered in (md, docx):
        assert "Acme ransomware" in rendered
        # The case has no IOCs, so the table renders its empty form rather than
        # being skipped - a missing annex is not the same as an empty one.
        assert "IOC" in rendered or "No IOC" in rendered
