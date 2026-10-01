"""
The annex tables are the document's design, and the TOC is a real Word field.

Remora painted every annex table itself: navy header, white bold text, grey
banding, 9/8 pt, with the colours as literals in the exporter. That is a visual
identity, and it was Remora's. A firm whose report template is set in their own
typeface and palette got four tables in ours in the middle of their deliverable,
and no amount of work in Word could change it.

These tests pin the convention that replaced it - a table style named
`Remora Annex`, then the document's own default, then the old rendering - and
they pin the half that is *not* delegated: column proportions, the `↳` that
marks a sub-technique, and the header row repeating across pages.

They also pin the table of contents as a Word field rather than a rendered list.
A list would be correct for exactly as long as nobody edited the document; the
field is what makes it clickable and what makes it survive an edit.
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient

from app.services import docx_style

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _upload(client: TestClient, name: str, body: bytes) -> int:
    response = client.post(
        "/api/v1/report-doc-templates/upload",
        data={"name": name, "description": ""},
        files={"file": (f"{name}.docx", body,
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    return response.json()["id"]


def _template_bytes(lines: list[str], annex_style_from: str | None = None) -> bytes:
    """
    A DOCX template.

    `annex_style_from` copies one of Word's built-in table styles under the
    name `Remora Annex`, which is how an author would create it: design a table
    in Word, save the look as a style, name it.
    """
    from docx import Document

    doc = Document()
    if annex_style_from:
        borrowed = doc.styles[annex_style_from]
        borrowed.element.set(f"{W}styleId", "RemoraAnnex")
        borrowed.element.find(f"{W}name").set(f"{W}val", docx_style.ANNEX_STYLE_NAME)
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _generate(client: TestClient, template_id: int, case_id: str):
    from docx import Document

    response = client.post(
        f"/api/v1/report-doc-templates/{template_id}/generate/{case_id}")
    assert response.status_code == 200, response.text
    return Document(io.BytesIO(response.content))


@pytest.fixture()
def case_with_annexes(db_session) -> str:
    """A case with enough in it that every annex table has rows."""
    import uuid

    from app.models.asset import Asset, AssetType
    from app.models.case import Case
    from app.models.ioc import IOC, IOCConfidence, IOCType

    case_id = str(uuid.uuid4())
    db_session.add(Case(
        id=case_id,
        title="Annex design",
        report_sections_data=json.dumps({"scope": "## Scope\n\n### Hosts\n\nTwo."}),
    ))
    db_session.flush()
    for i in range(3):
        db_session.add(IOC(case_id=case_id, type=IOCType.ip, value=f"10.0.0.{i}",
                           confidence=IOCConfidence.high, tlp="TLP:AMBER"))
    db_session.add(Asset(case_id=case_id, name="WKS-042", type=AssetType.workstation))
    db_session.commit()
    return case_id


def _first_table(doc):
    assert doc.tables, "the export produced no table"
    return doc.tables[0]


def _tbl_pr(table):
    return table._tbl.find(f"{W}tblPr")


# ─── Which design wins ────────────────────────────────────────────────────────

def test_a_plain_template_keeps_the_rendering_remora_always_had():
    """The fallback has to stay: most templates define no table style at all."""
    from docx import Document

    resolved = docx_style.annex_style(Document(io.BytesIO(_template_bytes(["x"]))))

    assert resolved.paints is True
    assert resolved.source == "built-in"


def test_a_remora_annex_style_in_the_template_wins():
    from docx import Document

    body = _template_bytes(["x"], annex_style_from="Light Grid Accent 1")
    resolved = docx_style.annex_style(Document(io.BytesIO(body)))

    assert resolved.name == docx_style.ANNEX_STYLE_NAME
    assert resolved.source == "template"
    assert resolved.paints is False


def test_words_own_borderless_default_is_not_a_choice():
    """
    `Normal Table` is the default in every document ever saved. Honouring it
    would mean every annex table losing its borders - an absent decision read
    as a decision for invisible tables.
    """
    from docx import Document

    resolved = docx_style.annex_style(Document(io.BytesIO(_template_bytes(["x"]))))

    assert resolved.name is None


def test_an_unreadable_document_does_not_fail_the_export():
    """A report template is a file somebody produced in some word processor."""
    class Hostile:
        @property
        def styles(self):
            raise RuntimeError("styles.xml is not what you think")

    assert docx_style.annex_style(Hostile()) is docx_style.BUILT_IN


# ─── What the style takes over, and what it does not ──────────────────────────

def test_remoras_painting_is_dropped_when_a_style_governs(auth_client, case_with_annexes):
    """
    The point of the whole convention: no Remora borders, no navy header fill,
    no banding. The style owns all three, and two sets of instructions on one
    table is how you get a navy header inside a client's grey design.
    """
    template = _upload(auth_client, "styled",
                       _template_bytes(["{{ioc_table}}"],
                                       annex_style_from="Light Grid Accent 1"))
    table = _first_table(_generate(auth_client, template, case_with_annexes))

    assert _tbl_pr(table).find(f"{W}tblBorders") is None, "Remora painted borders anyway"
    assert table._tbl.findall(f".//{W}shd") == [], "Remora shaded cells anyway"
    assert table.style.name == docx_style.ANNEX_STYLE_NAME


def test_remora_still_paints_when_the_template_asks_for_nothing(auth_client, case_with_annexes):
    template = _upload(auth_client, "plain", _template_bytes(["{{ioc_table}}"]))
    table = _first_table(_generate(auth_client, template, case_with_annexes))

    assert _tbl_pr(table).find(f"{W}tblBorders") is not None
    assert table._tbl.findall(f".//{W}shd"), "the header fill is gone"


def test_the_margin_bleed_belongs_to_remoras_design(auth_client, case_with_annexes):
    """
    The ~0.5 cm into each margin was chosen to go with Remora's palette. A
    client's table style expects its tables inside the text area, so the
    negative indent goes with the painting.
    """
    styled = _upload(auth_client, "styled-bleed",
                     _template_bytes(["{{asset_table}}"],
                                     annex_style_from="Light Grid Accent 1"))
    plain  = _upload(auth_client, "plain-bleed", _template_bytes(["{{asset_table}}"]))

    assert _tbl_pr(_first_table(_generate(auth_client, styled, case_with_annexes))) \
        .find(f"{W}tblInd") is None
    assert _tbl_pr(_first_table(_generate(auth_client, plain, case_with_annexes))) \
        .find(f"{W}tblInd") is not None


@pytest.mark.parametrize("annex_style_from", [None, "Light Grid Accent 1"])
def test_the_header_row_repeats_across_pages_either_way(
        auth_client, case_with_annexes, annex_style_from):
    """
    An annex of two hundred indicators runs over three pages, and on pages two
    and three the columns were unlabelled. That is not a look, it is a table
    that cannot be read.
    """
    template = _upload(auth_client, f"repeat-{annex_style_from}",
                       _template_bytes(["{{ioc_table}}"], annex_style_from))
    table = _first_table(_generate(auth_client, template, case_with_annexes))

    first_row = table.rows[0]._tr
    assert first_row.find(f"{W}trPr/{W}tblHeader") is not None


@pytest.mark.parametrize("annex_style_from", [None, "Light Grid Accent 1"])
def test_the_rows_are_the_same_whoever_paints_them(
        auth_client, case_with_annexes, annex_style_from):
    """Delegating the look must not delegate the content."""
    template = _upload(auth_client, f"rows-{annex_style_from}",
                       _template_bytes(["{{ioc_table}}"], annex_style_from))
    table = _first_table(_generate(auth_client, template, case_with_annexes))

    assert [c.text for c in table.rows[0].cells][:2] == ["Type", "Value"]
    assert len(table.rows) == 1 + 3
    assert "10.0.0.1" in {c.text for r in table.rows for c in r.cells}


def test_a_sub_technique_is_marked_whoever_paints_the_row(db_session):
    """
    The green tint went with Remora's palette; the arrow and the indent are what
    the row *means*. Under a client's style the tint goes and the arrow stays.
    """
    import uuid

    from docx import Document

    from app.models.case import Case
    from app.models.mitre import CaseTTP
    from app.routers.report_doc_templates import _build_mitre_word_table

    case_id = str(uuid.uuid4())
    case = Case(id=case_id, title="TTPs")
    db_session.add(case)
    db_session.flush()
    db_session.add(CaseTTP(case_id=case_id, tactic="execution", tactic_name="Execution",
                           technique_id="T1059.001", technique_name="PowerShell"))
    db_session.commit()
    db_session.refresh(case)

    doc   = Document()
    table = _build_mitre_word_table(
        doc, case, docx_style.AnnexStyle(name="Table Grid", source="template"))

    ids = [row.cells[1].text for row in table.rows]
    assert any("↳ T1059.001" in text for text in ids)
    assert table._tbl.findall(f".//{W}shd") == [], "the sub-technique tint survived"


# ─── The table of contents ────────────────────────────────────────────────────

def test_the_toc_is_a_word_field_not_a_list(auth_client, case_with_annexes):
    template = _upload(auth_client, "toc", _template_bytes(["{{toc}}", "{{scope}}"]))
    doc = _generate(auth_client, template, case_with_annexes)

    instructions = [el.text for el in doc.element.body.iter(f"{W}instrText")]
    assert any("TOC" in (text or "") for text in instructions), \
        "no TOC field instruction in the document"
    assert "{{toc}}" not in "\n".join(p.text for p in doc.paragraphs)


def test_the_toc_field_is_dirty_so_word_builds_it(auth_client, case_with_annexes):
    """
    Without this the field is there and empty, and the analyst has to know to
    press F9 on a document they are about to send to a client.
    """
    template = _upload(auth_client, "toc-dirty", _template_bytes(["{{toc}}"]))
    doc = _generate(auth_client, template, case_with_annexes)

    begins = [el for el in doc.element.body.iter(f"{W}fldChar")
              if el.get(f"{W}fldCharType") == "begin"]
    assert begins, "the field has no opening character"
    assert any(el.get(f"{W}dirty") == "true" for el in begins)


def test_the_toc_has_something_to_find(auth_client, case_with_annexes):
    """
    A TOC field reads heading styles. If the sections came out as plain
    paragraphs the field would build an empty table of contents and the failure
    would be invisible until a client opened the file.
    """
    template = _upload(auth_client, "toc-headings",
                       _template_bytes(["{{toc}}", "{{scope}}"]))
    doc = _generate(auth_client, template, case_with_annexes)

    styles = {p.style.name for p in doc.paragraphs if p.text.strip()}
    assert "Heading 2" in styles, f"no heading survived the export: {styles}"
    assert "Heading 3" in styles


def test_the_markdown_toc_lists_the_headings_of_the_finished_report(
        auth_client, case_with_annexes):
    """
    Markdown has no field to update, so a literal note saying "DOCX only" or an
    actual list were the options. The list is read from the rendered document,
    after the sections are in - the headings live inside them.
    """
    body = b"{{toc}}\n\n{{scope}}"
    response = auth_client.post(
        "/api/v1/report-doc-templates/upload",
        data={"name": "toc md", "description": ""},
        files={"file": ("toc.md", body, "text/markdown")},
    )
    template_id = response.json()["id"]

    text = auth_client.post(
        f"/api/v1/report-doc-templates/{template_id}/generate/{case_with_annexes}"
    ).content.decode()

    assert "- [Scope](#scope)" in text
    assert "  - [Hosts](#hosts)" in text, "the nesting of heading levels was lost"
    assert "{{toc}}" not in text


def test_the_markdown_toc_does_not_list_itself(auth_client, case_with_annexes):
    """It is built from the text with its own placeholder taken out."""
    rendered = docx_style.markdown_toc("## Scope\n\n- [Scope](#scope)")

    assert rendered.count("Scope") == 1


def test_a_heading_inside_a_code_block_is_not_a_heading():
    source = "# Real\n\n```\n## Not a heading\n```\n"

    assert docx_style.markdown_toc(source) == "- [Real](#real)"


def test_a_report_with_no_headings_says_so():
    assert "No headings" in docx_style.markdown_toc("Just a paragraph.")
